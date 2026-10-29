import io
import json
import time

import pandas as pd
import plotly.graph_objects as go
import streamlit as st
from google import genai
from google.genai import errors, types

# ----------------------------------------------------------------------------
# Page setup
# ----------------------------------------------------------------------------
st.set_page_config(page_title="Idea Spark AI", page_icon="✨", layout="wide")

# Models are tried in this order. If one is unavailable (404), rate-limited (429)
# or overloaded (503), the app moves on to the next one.
# Google renames models often: see https://ai.google.dev/gemini-api/docs/models
MODELS = [
    "gemini-3.6-flash",
    "gemini-3.5-flash",
    "gemini-3.1-flash-lite",
    "gemini-flash-latest",
]

MAX_FILE_MB = 15

# Speed setting -> how long the model "thinks" before answering
SPEEDS = {"⚡ Fast": "low", "⚖️ Balanced": "medium", "🧠 Deep thinking": "high"}
CHAT_SYSTEM = (
    "You are Spark, the friendly AI assistant inside the Idea Spark AI app. "
    "If someone asks your name, say you are Spark. If asked which model powers you, "
    "say Google Gemini. Be helpful, warm and clear."
)
MAX_HISTORY = 20  # only the last 20 chat messages are sent, which keeps replies quick

# ----------------------------------------------------------------------------
# Writing tools: name shown in the dropdown -> instruction sent to the AI
# Add your own by adding one more line.
# ----------------------------------------------------------------------------
TASKS = {
    "📝 Summarize": "Summarize the text in a few short, clear bullet points.",
    "✨ Rewrite professionally": "Rewrite the text in a polished, professional tone. Keep the meaning.",
    "🧒 Explain like I'm 12": "Explain the text in very simple words a 12-year-old can understand.",
    "✅ Fix grammar": "Fix all grammar, spelling and word-choice mistakes. Keep the meaning and style. "
                      "Return the corrected text first, then a short list of the main corrections.",
    "🎓 Explain at expert level (with references)": "Explain the topic of the text at an advanced, expert level. "
        "At the end, suggest real books, papers or organizations the reader can search for. "
        "Do NOT invent links or exact citations.",
    "🧠 Explain with analogies": "Explain the ideas in the text using clear everyday analogies.",
    "📚 Make study notes": "Turn the text into well-organized study notes with headings and bullet points.",
    "🃏 Make flashcards": "Create 8 flashcards from the text in the form 'Q: ... / A: ...'.",
    "🧩 Make a quiz (5 questions)": "Create a 5-question multiple-choice quiz from the text. "
                                    "Put the answer key at the end.",
    "🗺️ Mind-map outline": "Turn the text into a mind-map style nested outline.",
    "❓ Questions to think about": "Write 7 thought-provoking discussion questions based on the text.",
    "🤯 Interesting facts & trivia": "Share surprising, interesting facts and trivia related to the text's topic.",
    "📖 Turn into a short story": "Turn the ideas in the text into a short, engaging story.",
    "🎭 Turn into a poem": "Turn the ideas in the text into a poem.",
    "📱 Social media post": "Write a catchy social media post based on the text, with a few hashtags.",
    "💼 LinkedIn post": "Write a professional, engaging LinkedIn post based on the text.",
    "📧 Draft an email": "Draft a clear, polite email based on the text, with a subject line.",
    "🎤 Presentation script": "Write a short presentation script (about 2 minutes) based on the text.",
    "🏷️ Title & headline ideas": "Suggest 10 creative titles or headlines for the text.",
    "🎯 Key takeaways & action items": "List the key takeaways and clear action items from the text.",
    "⚖️ Pros and cons": "List the pros and cons related to the text in two clear sections.",
    "🔍 Critique & counter-arguments": "Critique the text fairly and give strong counter-arguments.",
    "🌍 Translate": "Translate the text into {language}. Return only the translation.",
}

# ----------------------------------------------------------------------------
# Light, colorful styling
# ----------------------------------------------------------------------------
# Cute background doodles (inline SVG, no image files needed)
DOODLE_SVG = """<svg xmlns='http://www.w3.org/2000/svg' width='440' height='440' viewBox='0 0 440 440' fill='none' stroke-linecap='round' stroke-linejoin='round'>
<g stroke='%23B9ABF7' stroke-width='2.5' opacity='.55'>
<path d='M50 50 q8 -18 16 0 q8 18 16 0 q8 -18 16 0'/>
<circle cx='330' cy='70' r='14' stroke-dasharray='4 7'/>
<path d='M205 125 v22 M194 136 h22'/>
<path d='M60 380 q10 -14 20 0 q10 14 20 0'/>
</g>
<g stroke='%23F6A6D2' stroke-width='2.5' opacity='.6'>
<path transform='translate(360 270) scale(.7)' d='M0 -4 C-14 -22 -34 -2 0 24 C34 -2 14 -22 0 -4z'/>
<path d='M150 250 l5 -11 l5 11 l11 5 l-11 5 l-5 11 l-5 -11 l-11 -5z'/>
<path d='M300 380 l3 -7 l3 7 l7 3 l-7 3 l-3 7 l-3 -7 l-7 -3z'/>
</g>
<g stroke='%23FFBE94' stroke-width='2.5' opacity='.65'>
<circle cx='95' cy='200' r='15'/>
<path d='M88 218 h14 M90 225 h10'/>
<path d='M95 170 v-8 M72 178 l-6 -5 M118 178 l6 -5'/>
<path d='M250 40 q10 -10 20 0 q10 10 20 0'/>
</g>
<g stroke='%239ADBC5' stroke-width='2.5' opacity='.65'>
<path d='M380 150 l4 -9 l4 9 l9 4 l-9 4 l-4 9 l-4 -9 l-9 -4z'/>
<circle cx='230' cy='330' r='7'/>
<path d='M390 390 v16 M382 398 h16'/>
</g>
<g fill='%23D9CFFF' opacity='.7'><circle cx='180' cy='60' r='4'/><circle cx='20' cy='290' r='4'/><circle cx='420' cy='330' r='4'/></g>
<g fill='%23FFD3E8' opacity='.8'><circle cx='320' cy='190' r='4'/><circle cx='120' cy='120' r='4'/></g>
</svg>"""
DOODLE_URI = 'url("data:image/svg+xml;utf8,' + DOODLE_SVG.replace("\n", "") + '")'

CSS = """
<style>
.stApp {
    background-image: __DOODLE__, linear-gradient(135deg, #F4EFFF 0%, #FFF1F8 55%, #FFF5EC 100%);
    background-size: 440px 440px, cover;
    background-repeat: repeat, no-repeat;
    background-attachment: fixed, fixed;
}
.stApp, .stApp p, .stApp li, .stApp label, .stApp span, .stApp div,
.stApp h2, .stApp h3, .stApp h4 { color: #2D2A4A; }
[data-testid="stSidebar"] { background: linear-gradient(180deg, rgba(244,239,255,.95), rgba(255,240,248,.95)); border-right: 1px solid #E6DEFF; }
[data-testid="stSidebar"] > div { background: transparent; }
[data-testid="stSidebar"] * { color: #2D2A4A; }
[data-testid="stHeader"] { background: transparent; }

.hero { padding: 6px 0 14px 0; }
.hero h1 { font-size: 2.4rem; margin: 0; background: linear-gradient(90deg, #7C6AF0, #E879B9);
    -webkit-background-clip: text; background-clip: text; color: transparent !important;
    -webkit-text-fill-color: transparent; }
.hero p { margin: 4px 0 0 0; color: #6B6890 !important; }

button[data-baseweb="tab"] { background: transparent; font-weight: 600; color: #7A7699 !important; }
button[data-baseweb="tab"][aria-selected="true"] { color: #7C6AF0 !important; }

[data-testid="stChatMessage"] {
    background: rgba(255,255,255,.82); border: 1px solid #EADFFF; border-radius: 18px;
    padding: 12px 16px; box-shadow: 0 4px 14px rgba(139, 124, 246, 0.08);
}
.card { background: rgba(255,255,255,.82); border: 1px solid #EADFFF; border-radius: 18px;
        padding: 18px; box-shadow: 0 4px 14px rgba(139, 124, 246, 0.08); }

.stTextArea textarea, .stTextInput input {
    background: rgba(255,255,255,.85) !important; color: #2D2A4A !important;
    border: 1px solid #DDD3FF !important; border-radius: 14px !important;
}
div[data-baseweb="select"] > div { background: rgba(255,255,255,.85) !important; border-radius: 14px !important; border-color: #DDD3FF !important; }
div[data-baseweb="select"] * { color: #2D2A4A !important; }
ul[role="listbox"], div[data-baseweb="popover"] > div { background: #FFFFFF !important; }
ul[role="listbox"] li { color: #2D2A4A !important; }
ul[role="listbox"] li:hover { background: #F1ECFF !important; }

/* Chat bar: soft glass instead of a big white strip */
[data-testid="stBottom"], [data-testid="stBottom"] > div,
[data-testid="stBottomBlockContainer"] { background: transparent !important; }
[data-testid="stChatInput"] { background: rgba(255,255,255,.78) !important; border: 1px solid #DDD3FF !important;
    border-radius: 20px !important; box-shadow: 0 6px 20px rgba(139,124,246,.14); }
[data-testid="stChatInput"] > div, [data-testid="stChatInput"] textarea { background: transparent !important; color: #2D2A4A !important; }
[data-testid="stChatInput"] button { background: linear-gradient(135deg, #8B7CF6, #F08CC8) !important; border-radius: 12px; }
[data-testid="stChatInput"] button svg { fill: #FFFFFF !important; color: #FFFFFF !important; }

[data-testid="stFileUploaderDropzone"] { background: rgba(255,255,255,.7); border: 2px dashed #CFC4FF; border-radius: 16px; }
[data-testid="stFileUploaderDropzone"] * { color: #2D2A4A !important; }

.stButton > button, [data-testid="stFileUploaderDropzone"] button {
    background: linear-gradient(90deg, #8B7CF6, #F08CC8); color: #FFFFFF !important;
    border: none; border-radius: 14px; padding: 0.55rem 1.3rem; font-weight: 600;
    box-shadow: 0 4px 12px rgba(139,124,246,.25);
}
.stButton > button:hover { filter: brightness(1.06); }
.stButton > button p, [data-testid="stFileUploaderDropzone"] button * { color: #FFFFFF !important; }

.thinking { display: inline-flex; align-items: center; gap: 10px; padding: 8px 16px;
    background: rgba(255,255,255,.85); border: 1px solid #DDD3FF; border-radius: 999px;
    color: #7C6AF0; font-weight: 600; animation: glow 1.8s ease-in-out infinite; }
.thinking .dot { width: 8px; height: 8px; border-radius: 50%; background: #8B7CF6;
    animation: bounce 1.2s infinite ease-in-out; }
.thinking .dot:nth-child(2) { animation-delay: 0.15s; background: #F08CC8; }
.thinking .dot:nth-child(3) { animation-delay: 0.3s; background: #FFB27D; }
@keyframes bounce { 0%, 80%, 100% { transform: translateY(0); opacity: .5; }
                    40% { transform: translateY(-6px); opacity: 1; } }
@keyframes glow { 0%, 100% { box-shadow: 0 0 0 0 rgba(139,124,246,0.25); }
                  50% { box-shadow: 0 0 0 8px rgba(139,124,246,0); } }

/* Always light: no black/dark parts, whatever the browser or system theme is */
:root, html, body, .stApp { color-scheme: light !important; }
[data-testid="stMainMenuPopover"], [data-testid="stMainMenuList"],
[data-testid="stMainMenuPopover"] * { background: #FFFFFF !important; color: #2D2A4A !important; }
[data-testid="stToolbar"] *, [data-testid="stHeader"] button { color: #2D2A4A !important; }
[data-testid="stTooltipContent"] { background: #FFFFFF !important; color: #2D2A4A !important; }
[data-testid="stCode"] pre { background: #F4F0FF !important; }
[data-testid="stCode"] code, [data-testid="stCode"] code span { color: #2D2A4A !important; }
[data-testid="stExpander"] { background: rgba(255,255,255,.7); border: 1px solid #EADFFF; border-radius: 14px; }
[data-testid="stExpander"] summary, [data-testid="stExpander"] summary * { color: #2D2A4A !important; }
[data-testid="stDataFrame"] { border: 1px solid #EADFFF; border-radius: 12px; overflow: hidden; }
.stDownloadButton > button { background: #FFFFFF; color: #7C6AF0 !important; border: 1px solid #DDD3FF;
    border-radius: 14px; font-weight: 600; }
.stDownloadButton > button p { color: #7C6AF0 !important; }
</style>
"""
st.markdown(CSS.replace("__DOODLE__", DOODLE_URI), unsafe_allow_html=True)


def pill(text: str) -> str:
    return (
        f'<div class="thinking"><span>{text}</span>'
        '<span class="dot"></span><span class="dot"></span><span class="dot"></span></div>'
    )


# ----------------------------------------------------------------------------
# Gemini helpers
# ----------------------------------------------------------------------------
class FriendlyError(Exception):
    """An error whose message is safe to show to the user."""

    def __init__(self, message, detail=""):
        super().__init__(message)
        self.detail = detail


def make_config(temperature, system, thinking, json_mode=False):
    kwargs = dict(temperature=temperature, system_instruction=system)
    if json_mode:
        kwargs["response_mime_type"] = "application/json"
    if thinking:
        kwargs["thinking_config"] = types.ThinkingConfig(thinking_level=thinking)
    return types.GenerateContentConfig(**kwargs)


ROUNDS = 5  # how many times all models are tried before giving up


def stream_gemini(client, contents, temperature, model_list, system=None,
                  thinking="low", on_wait=None, json_mode=False):
    """Yield text chunks without bothering the user.

    Busy or rate-limited? Jump to the next model right away. If every model is
    busy, wait a little (longer each round) and go around again.
    """
    config = make_config(temperature, system, thinking, json_mode)
    dead = set()  # models that do not exist (404) are skipped for this request
    last_detail = ""
    final_problems = set()

    for round_no in range(ROUNDS):
        live = [m for m in model_list if m not in dead]
        if not live:
            break
        problems = set()

        for model in live:
            for _ in range(2):  # 2nd pass only if we must drop the speed setting
                got_text = False
                try:
                    for chunk in client.models.generate_content_stream(
                        model=model, contents=contents, config=config
                    ):
                        if chunk.text:
                            got_text = True
                            yield chunk.text
                    return
                except errors.APIError as e:
                    code = getattr(e, "code", None)
                    last_detail = f"{model} -> {code}: {str(e)[:300]}"
                    if got_text:
                        raise FriendlyError("The answer was cut off. Please try again.", last_detail)
                    if code == 400 and "thinking" in str(e).lower() and config.thinking_config:
                        config = make_config(temperature, system, None, json_mode)
                        continue  # same model, without the speed setting
                    if code in (400, 401, 403):
                        if code in (401, 403) or "api key" in str(e).lower():
                            raise FriendlyError("Your API key was rejected. Please check it in the sidebar.", last_detail)
                        raise FriendlyError("Google could not process this request.", last_detail)
                    if code == 404:
                        dead.add(model)
                    elif code == 429:
                        problems.add("limit")
                    elif code in (500, 502, 503, 504):
                        problems.add("busy")
                    else:
                        raise FriendlyError("Unexpected error from Google.", last_detail)
                    break  # next model

        final_problems = problems
        if not problems or round_no == ROUNDS - 1:
            break
        if problems == {"limit"}:
            wait, msg = 5 * (round_no + 1), "⏳ Free-tier limit reached, waiting a moment"
        else:
            wait, msg = 2 * (round_no + 1), "🔄 Google is busy, trying again"
        if on_wait:
            on_wait(msg)
        time.sleep(wait)

    if "limit" in final_problems:
        raise FriendlyError("Free-tier limit reached. Please wait about a minute and send again.", last_detail)
    if "busy" in final_problems:
        raise FriendlyError("Google's servers are still busy. Please send your message again in a moment.", last_detail)
    raise FriendlyError(
        "None of the models were available. Type a current model name in the sidebar "
        "(see ai.google.dev/gemini-api/docs/models).", last_detail,
    )


def run_stream(client, contents, temperature, model_list, thinking_text,
               system=None, thinking="low"):
    """Show the thinking pill, stream the answer, return the final text (or None on error)."""
    pill_box = st.empty()
    out_box = st.empty()
    detail_box = st.empty()
    pill_box.markdown(pill(thinking_text), unsafe_allow_html=True)

    def on_wait(msg):
        pill_box.markdown(pill(msg), unsafe_allow_html=True)

    text = ""
    try:
        for piece in stream_gemini(client, contents, temperature, model_list,
                                   system, thinking, on_wait):
            pill_box.empty()
            text += piece
            out_box.markdown(text + "▌")
        pill_box.empty()
        out_box.markdown(text)
        return text
    except FriendlyError as e:
        pill_box.empty()
        out_box.warning(str(e))
        if e.detail:
            with detail_box.expander("Technical details"):
                st.code(e.detail)
        return None
    except Exception as e:  # last-resort safety net
        pill_box.empty()
        out_box.error(f"Something went wrong: {e}")
        return None


def read_upload(uploaded):
    """Return (text_or_None, Part_or_None, error_or_None) for an uploaded file."""
    data = uploaded.getvalue()
    if not data:
        return None, None, "That file is empty."
    if len(data) > MAX_FILE_MB * 1024 * 1024:
        return None, None, f"That file is larger than {MAX_FILE_MB} MB."

    name = uploaded.name.lower()
    try:
        if name.endswith(".pdf"):
            return None, types.Part.from_bytes(data=data, mime_type="application/pdf"), None
        if name.endswith((".png", ".jpg", ".jpeg")):
            mime = "image/png" if name.endswith(".png") else "image/jpeg"
            return None, types.Part.from_bytes(data=data, mime_type=mime), None
        if name.endswith(".docx"):
            from docx import Document

            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            text = "\n".join(parts).strip()
            return (text, None, None) if text else (None, None, "No readable text in that Word file.")
        text = data.decode("utf-8", errors="ignore").strip()
        return (text, None, None) if text else (None, None, "No readable text in that file.")
    except Exception as e:
        return None, None, f"Could not read that file: {e}"


# ----------------------------------------------------------------------------
# Graph helpers
# ----------------------------------------------------------------------------
CHART_TYPES = {
    "✨ Auto (AI chooses)": None,
    "📊 Bar": "bar",
    "📏 Horizontal bar": "horizontal_bar",
    "📈 Line": "line",
    "🏔️ Area": "area",
    "🥧 Pie": "pie",
    "🔵 Scatter": "scatter",
    "📶 Histogram": "histogram",
}
VALID_KINDS = {"bar", "horizontal_bar", "line", "area", "pie", "scatter", "histogram"}
PALETTE = ["#8B7CF6", "#F08CC8", "#FFB27D", "#7DD3B0", "#6CB4F5", "#F5D76E"]

GRAPH_SYSTEM = """You turn ideas, notes, text and tables into chart data.
Reply with ONLY one valid JSON object (no markdown, no explanation) with these keys:
{"title": str, "chart_type": "bar"|"horizontal_bar"|"line"|"area"|"pie"|"scatter"|"histogram",
 "x_label": str, "y_label": str, "categories": [labels or x numbers],
 "series": [{"name": str, "values": [numbers]}], "summary": "1-2 sentence takeaway",
 "estimated": true|false}
Rules: at most 30 categories and 5 series. Every series must have exactly as many values as
categories (for a histogram, categories can be empty and the values are the raw numbers).
Values must be plain numbers (no units, no % signs). Pick the chart type that fits the data best
unless one is requested. If the user gives no numbers, create reasonable illustrative estimates
and set estimated to true. Otherwise use the given numbers exactly and set estimated to false."""


def to_num(v):
    try:
        return float(str(v).replace(",", "").replace("%", "").strip())
    except (ValueError, TypeError):
        return None


def parse_json(text):
    """Parse the model's JSON, tolerating code fences or extra words around it."""
    text = text.strip()
    try:
        return json.loads(text)
    except ValueError:
        pass
    a, b = text.find("{"), text.rfind("}")
    if a != -1 and b > a:
        return json.loads(text[a:b + 1])
    raise ValueError("No JSON found")


def normalize_spec(raw, forced_kind=None):
    """Clean up the AI's chart data so it is always safe to draw."""
    if not isinstance(raw, dict):
        raise ValueError("Chart data was not an object")
    kind = forced_kind or raw.get("chart_type")
    if kind not in VALID_KINDS:
        kind = "bar"

    cats = list(raw.get("categories") or [])[:30]
    series = []
    for s_ in (raw.get("series") or [])[:5]:
        if not isinstance(s_, dict):
            continue
        vals = [to_num(v) for v in (s_.get("values") or [])]
        series.append({"name": str(s_.get("name") or "Value"), "values": vals})
    series = [s_ for s_ in series if any(v is not None for v in s_["values"])]
    if not series:
        raise ValueError("No numbers to plot")

    if kind != "histogram":
        n = len(cats) if cats else max(len(s_["values"]) for s_ in series)
        if not cats:
            cats = [str(i + 1) for i in range(n)]
        for s_ in series:  # shorter series get empty gaps instead of cutting data off
            s_["values"] = (s_["values"] + [None] * n)[:n]
        if kind == "scatter":
            nums = [to_num(c) for c in cats]
            cats = nums if all(v is not None for v in nums) else list(range(1, n + 1))
        else:
            cats = [str(c) for c in cats]

    return {
        "title": str(raw.get("title") or "Your graph"),
        "chart_type": kind,
        "x_label": str(raw.get("x_label") or ""),
        "y_label": str(raw.get("y_label") or ""),
        "categories": cats,
        "series": series,
        "summary": str(raw.get("summary") or ""),
        "estimated": bool(raw.get("estimated")),
    }


def build_figure(spec):
    kind, cats, series = spec["chart_type"], spec["categories"], spec["series"]
    fig = go.Figure()
    for i, s_ in enumerate(series):
        color = PALETTE[i % len(PALETTE)]
        if kind == "pie":
            fig.add_trace(go.Pie(labels=cats, values=s_["values"], hole=0.35, name=s_["name"],
                                 marker=dict(colors=PALETTE), textinfo="label+percent"))
            break
        if kind == "histogram":
            fig.add_trace(go.Histogram(x=s_["values"], name=s_["name"], marker_color=color, opacity=0.85))
        elif kind == "scatter":
            fig.add_trace(go.Scatter(x=cats, y=s_["values"], name=s_["name"], mode="markers",
                                     marker=dict(size=12, color=color, line=dict(width=1, color="white"))))
        elif kind == "line":
            fig.add_trace(go.Scatter(x=cats, y=s_["values"], name=s_["name"], mode="lines+markers",
                                     line=dict(width=3, color=color), marker=dict(size=8)))
        elif kind == "area":
            fig.add_trace(go.Scatter(x=cats, y=s_["values"], name=s_["name"], mode="lines",
                                     line=dict(width=2, color=color), fill="tozeroy"))
        elif kind == "horizontal_bar":
            fig.add_trace(go.Bar(x=s_["values"], y=cats, name=s_["name"], orientation="h", marker_color=color))
        else:
            fig.add_trace(go.Bar(x=cats, y=s_["values"], name=s_["name"], marker_color=color))

    x_title, y_title = spec["x_label"], spec["y_label"]
    if kind == "horizontal_bar":
        x_title, y_title = y_title, x_title
    fig.update_layout(
        title=dict(text=spec["title"], font=dict(size=20)),
        template="plotly_white", barmode="group", colorway=PALETTE,
        font=dict(color="#2D2A4A"),
        paper_bgcolor="rgba(255,255,255,0)", plot_bgcolor="rgba(255,255,255,0.6)",
        xaxis_title=x_title, yaxis_title=y_title,
        margin=dict(l=40, r=20, t=70, b=40), legend=dict(orientation="h", y=-0.2),
    )
    return fig


def spec_to_frame(spec):
    if spec["chart_type"] == "histogram":
        return pd.DataFrame({s_["name"]: pd.Series(s_["values"]) for s_ in spec["series"]})
    data = {"label": spec["categories"]}
    for s_ in spec["series"]:
        data[s_["name"]] = s_["values"]
    return pd.DataFrame(data)


def read_table(uploaded):
    """Read a CSV or Excel upload into a small DataFrame (None, error on failure)."""
    data = uploaded.getvalue()
    if not data:
        return None, "That file is empty."
    if len(data) > MAX_FILE_MB * 1024 * 1024:
        return None, f"That file is larger than {MAX_FILE_MB} MB."
    try:
        if uploaded.name.lower().endswith(".xlsx"):
            df = pd.read_excel(io.BytesIO(data))
        else:
            df = pd.read_csv(io.BytesIO(data))
    except Exception as e:
        return None, f"Could not read that file: {e}"
    if df.empty:
        return None, "No rows found in that file."
    return df.iloc[:200, :12], None


# ----------------------------------------------------------------------------
# Sidebar
# ----------------------------------------------------------------------------
with st.sidebar:
    st.header("⚙️ Settings")

    api_key = ""
    try:
        api_key = st.secrets.get("GEMINI_API_KEY", "")
    except Exception:
        pass
    if not api_key:
        api_key = st.text_input("Gemini API key", type="password",
                                help="Free key at aistudio.google.com")

    model_choice = st.selectbox("Model", MODELS)
    custom_model = st.text_input("Custom model name (optional)",
                                 placeholder="e.g. gemini-3.6-flash")
    temperature = st.slider("Creativity", 0.0, 2.0, 1.0, 0.1)
    speed = st.selectbox(
        "Speed",
        list(SPEEDS.keys()),
        help="Fast answers quickly. Deep thinking is slower but better for hard questions.",
    )
    thinking_level = SPEEDS[speed]

    if st.button("🗑️ Clear chat"):
        st.session_state.messages = []
        st.rerun()

    st.caption("Get a free key at aistudio.google.com. Never share it publicly.")

# Chosen model first, then the rest as backups
first = custom_model.strip() or model_choice
model_list = [first] + [m for m in MODELS if m != first]

# ----------------------------------------------------------------------------
# Header
# ----------------------------------------------------------------------------
st.markdown(
    '<div class="hero"><h1>✨ Idea Spark AI</h1>'
    "<p>Spark new ideas, chat with AI and use handy writing tools, powered by Gemini Flash.</p></div>",
    unsafe_allow_html=True,
)

if "messages" not in st.session_state:
    st.session_state.messages = []

@st.cache_resource
def get_client(key):
    return genai.Client(api_key=key)


client = get_client(api_key) if api_key else None
tab_chat, tab_tools, tab_graph = st.tabs(["💬 Chat", "🛠️ Writing Tools", "📊 Graphs"])

# ----------------------------------------------------------------------------
# Chat tab
# ----------------------------------------------------------------------------
with tab_chat:
    chat_area = st.container()  # messages live here, above the input box
    prompt = st.chat_input("Ask me anything...")

    with chat_area:
        for m in st.session_state.messages:
            with st.chat_message(m["role"]):
                st.markdown(m["content"])

        if prompt:
            if not client:
                st.info("🔑 Add your Gemini API key in the sidebar to start.")
            else:
                st.session_state.messages.append({"role": "user", "content": prompt})
                with st.chat_message("user"):
                    st.markdown(prompt)

                history = [
                    types.Content(
                        role="user" if m["role"] == "user" else "model",
                        parts=[types.Part.from_text(text=m["content"])],
                    )
                    for m in st.session_state.messages[-MAX_HISTORY:]
                ]
                with st.chat_message("assistant"):
                    reply = run_stream(client, history, temperature, model_list,
                                       "✨ Spark is thinking", system=CHAT_SYSTEM,
                                       thinking=thinking_level)
                if reply:
                    st.session_state.messages.append({"role": "assistant", "content": reply})
                else:
                    st.session_state.messages.pop()  # drop the unanswered question

# ----------------------------------------------------------------------------
# Writing tools tab
# ----------------------------------------------------------------------------
with tab_tools:
    st.markdown("#### Pick a tool, add your text or file, and press Run")
    task_name = st.selectbox("What would you like to do?", list(TASKS.keys()))

    language = ""
    if task_name == "🌍 Translate":
        language = st.text_input("Translate into which language?", placeholder="e.g. Urdu, French")

    if "Expert level" in task_name or "expert level" in task_name:
        st.caption("⚠️ AI can get references wrong. Always verify them before relying on them.")

    source = st.radio("Input", ["✍️ Paste text", "📎 Upload file"], horizontal=True,
                      label_visibility="collapsed")

    text_in, file_in = "", None
    if source == "✍️ Paste text":
        text_in = st.text_area("Your text", height=220, placeholder="Paste your text here...")
    else:
        file_in = st.file_uploader(
            f"Upload a file (up to {MAX_FILE_MB} MB)",
            type=["pdf", "docx", "txt", "md", "csv", "png", "jpg", "jpeg"],
        )

    if st.button("🚀 Run"):
        if not client:
            st.info("🔑 Add your Gemini API key in the sidebar to start.")
        elif task_name == "🌍 Translate" and not language.strip():
            st.warning("Please type the language you want to translate into.")
        else:
            instruction = TASKS[task_name].format(language=language.strip() or "English")
            contents = None

            if source == "✍️ Paste text":
                if not text_in.strip():
                    st.warning("Please paste some text first.")
                else:
                    contents = [types.Part.from_text(text=f"{instruction}\n\nText:\n{text_in}")]
            else:
                if not file_in:
                    st.warning("Please upload a file first.")
                else:
                    file_text, file_part, err = read_upload(file_in)
                    if err:
                        st.warning(err)
                    elif file_part is not None:
                        contents = [types.Part.from_text(text=instruction + "\n\nThe content is in the attached file."),
                                    file_part]
                    else:
                        contents = [types.Part.from_text(text=f"{instruction}\n\nText:\n{file_text}")]

            if contents:
                st.markdown('<div class="card">', unsafe_allow_html=True)
                run_stream(
                    client, contents, temperature, model_list, "✨ Crafting your result",
                    system="Follow the task exactly. Output only the result, formatted in clean Markdown.",
                    thinking=thinking_level,
                )
                st.markdown("</div>", unsafe_allow_html=True)


# ----------------------------------------------------------------------------
# Graphs tab
# ----------------------------------------------------------------------------
with tab_graph:
    st.markdown("#### Turn any idea or information into a graph")
    g_text = st.text_area(
        "Describe your idea or paste your information",
        height=150,
        placeholder="e.g. Compare the monthly sales of 3 products: A 120, 150, 170 / B 90, 110, 95 / C 60, 80, 130 "
                    "(Jan to Mar)\nor: Show how a student could split a 24-hour day.",
    )
    g_kind_name = st.selectbox("Chart type", list(CHART_TYPES.keys()))
    with st.expander("📎 Or use a CSV / Excel file"):
        g_file = st.file_uploader(f"Upload a CSV or Excel file (up to {MAX_FILE_MB} MB)",
                                  type=["csv", "xlsx"], key="graph_file")

    if st.button("📊 Create graph"):
        if not client:
            st.info("🔑 Add your Gemini API key in the sidebar to start.")
        elif not g_text.strip() and not g_file:
            st.warning("Please describe your idea, paste some information, or upload a file.")
        else:
            prompt_parts = [g_text.strip() or "Chart the most interesting insight in this data."]
            problem = None
            if g_file:
                table, problem = read_table(g_file)
                if problem is None:
                    prompt_parts.append("Data (CSV):\n" + table.to_csv(index=False))
            forced = CHART_TYPES[g_kind_name]
            if forced:
                prompt_parts.append(f"Use a {forced.replace('_', ' ')} chart.")

            if problem:
                st.warning(problem)
            else:
                box = st.empty()
                box.markdown(pill("📊 Drawing your graph"), unsafe_allow_html=True)
                try:
                    raw_text = "".join(stream_gemini(
                        client, [types.Part.from_text(text="\n\n".join(prompt_parts))],
                        0.4, model_list, system=GRAPH_SYSTEM, thinking=thinking_level,
                        on_wait=lambda msg: box.markdown(pill(msg), unsafe_allow_html=True),
                        json_mode=True,
                    ))
                    st.session_state.graph = normalize_spec(parse_json(raw_text), forced)
                    box.empty()
                except FriendlyError as e:
                    box.empty()
                    st.warning(str(e))
                    if e.detail:
                        with st.expander("Technical details"):
                            st.code(e.detail)
                except (ValueError, KeyError, TypeError):
                    box.empty()
                    st.warning("I could not turn that into a graph. Try adding a few numbers or "
                               "describing what you want to compare.")

    spec = st.session_state.get("graph")
    if spec:
        st.plotly_chart(build_figure(spec), use_container_width=True)
        if spec["summary"]:
            st.markdown(f"💡 **{spec['summary']}**")
        if spec["estimated"]:
            st.caption("⚠️ These numbers are AI estimates, so treat them as an illustration, not facts.")
        frame = spec_to_frame(spec)
        with st.expander("🔢 See the data"):
            st.dataframe(frame, use_container_width=True)
        d1, d2 = st.columns(2)
        with d1:
            st.download_button("⬇️ Download data (CSV)", frame.to_csv(index=False),
                               file_name="idea-spark-data.csv", mime="text/csv")
        with d2:
            st.download_button("⬇️ Download graph (HTML)",
                               build_figure(spec).to_html(include_plotlyjs="cdn"),
                               file_name="idea-spark-graph.html", mime="text/html")
        st.caption("Tip: hover over the graph and click the camera icon to save it as a PNG image.")
